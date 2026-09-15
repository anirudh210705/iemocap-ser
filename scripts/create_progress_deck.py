from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("BTP_SER_Progress_Update.pptx")
NAVY = RGBColor(15, 31, 55)
BLUE = RGBColor(39, 110, 241)
CYAN = RGBColor(40, 194, 210)
GREEN = RGBColor(38, 166, 91)
AMBER = RGBColor(245, 166, 35)
LIGHT = RGBColor(244, 247, 252)
MID = RGBColor(99, 115, 135)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(230, 239, 255)
PALE_GREEN = RGBColor(230, 247, 237)
PALE_AMBER = RGBColor(255, 245, 221)


def add_text(slide, text, x, y, w, h, size=18, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def circle(slide, x, y, d, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    return shape


def base_slide(prs, number, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE
    rect(slide, 0, 0, 13.333, 0.12, BLUE, radius=False)
    add_text(slide, f"0{number}", 0.55, 0.42, 0.55, 0.35, 11, BLUE, True)
    add_text(slide, title, 1.15, 0.3, 11.2, 0.65, 27, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 1.15, 0.9, 11.0, 0.42, 12, MID)
    add_text(slide, "BTP • Speech Emotion Recognition", 0.58, 7.12, 4.0, 0.22, 9, MID)
    add_text(slide, str(number), 12.15, 7.08, 0.55, 0.25, 9, MID, align=PP_ALIGN.RIGHT)
    return slide


def add_step(slide, x, y, w, title, detail, accent=BLUE):
    rect(slide, x, y, w, 1.2, LIGHT)
    rect(slide, x, y, 0.09, 1.2, accent, radius=False)
    add_text(slide, title, x + 0.25, y + 0.15, w - 0.4, 0.32, 15, NAVY, True)
    add_text(slide, detail, x + 0.25, y + 0.53, w - 0.4, 0.48, 11, MID)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.16, 7.5, CYAN, radius=False)
    circle(slide, 10.55, 0.6, 1.85, BLUE)
    circle(slide, 11.5, 1.82, 0.78, CYAN)
    for i, height in enumerate((0.45, 0.85, 1.35, 0.72, 1.12, 0.55, 0.95)):
        rect(slide, 9.2 + i * 0.42, 4.9 - height / 2, 0.18, height, CYAN if i % 2 else BLUE, radius=True)
    add_text(slide, "BTP PROGRESS UPDATE", 0.9, 0.75, 5.5, 0.4, 13, CYAN, True)
    add_text(slide, "Speech Emotion\nRecognition", 0.9, 1.35, 7.8, 1.8, 38, WHITE, True)
    add_text(slide, "Speaker-independent emotion classification on IEMOCAP\nusing CNN and pretrained speech encoders", 0.95, 3.45, 7.5, 1.0, 18, RGBColor(200, 211, 227))
    rect(slide, 0.95, 5.4, 3.05, 0.68, BLUE)
    add_text(slide, "CURRENT STAGE  •  CODE READY", 1.13, 5.55, 2.7, 0.3, 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "August 2026", 0.95, 6.58, 2.2, 0.3, 11, RGBColor(170, 184, 204))

    # Slide 2
    slide = base_slide(prs, 2, "What exactly am I doing?", "A reproducible study of speech emotion recognition under speaker-independent evaluation")
    steps = [
        ("Speech input", "16 kHz IEMOCAP\nutterance", BLUE),
        ("Representation", "Log-Mel CNN or\nWav2Vec2/HuBERT", CYAN),
        ("Pooling", "Attentive temporal\nsummarization", AMBER),
        ("Prediction", "Angry • Happy\nNeutral • Sad", GREEN),
    ]
    for i, (title, detail, color) in enumerate(steps):
        x = 0.75 + i * 3.05
        circle(slide, x + 0.83, 1.65, 0.72, color)
        add_text(slide, str(i + 1), x + 0.83, 1.65, 0.72, 0.72, 18, WHITE, True, PP_ALIGN.CENTER)
        rect(slide, x, 2.55, 2.4, 1.35, LIGHT)
        add_text(slide, title, x + 0.18, 2.73, 2.04, 0.32, 15, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.18, 3.13, 2.04, 0.5, 11, MID, False, PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "→", x + 2.48, 2.93, 0.5, 0.42, 24, BLUE, True, PP_ALIGN.CENTER)
    rect(slide, 0.85, 4.55, 11.65, 1.4, PALE_BLUE)
    add_text(slide, "Research focus", 1.12, 4.78, 1.8, 0.35, 14, BLUE, True)
    add_text(slide, "Compare a conventional CNN baseline with pretrained speech encoders, then measure how augmentation affects macro-F1 without allowing speaker leakage.", 2.85, 4.68, 9.15, 0.72, 16, NAVY)

    # Slide 3
    slide = base_slide(prs, 3, "Dataset prepared and validated", "Official IEMOCAP annotations mapped to a standard four-emotion benchmark")
    metrics = [("5,531", "usable utterances"), ("10", "unique speakers"), ("5", "recording sessions"), ("4", "emotion classes")]
    for i, (value, label) in enumerate(metrics):
        x = 0.75 + i * 3.05
        rect(slide, x, 1.48, 2.55, 1.18, LIGHT)
        add_text(slide, value, x + 0.15, 1.62, 2.25, 0.48, 27, BLUE, True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.15, 2.13, 2.25, 0.25, 11, MID, False, PP_ALIGN.CENTER)
    labels = [("Neutral", 1708, RGBColor(64, 116, 211)), ("Happy + Excited", 1636, CYAN), ("Angry", 1103, AMBER), ("Sad", 1084, RGBColor(130, 108, 190))]
    max_count = 1800
    add_text(slide, "Class distribution", 0.85, 3.12, 2.4, 0.38, 15, NAVY, True)
    for i, (label, count, color) in enumerate(labels):
        y = 3.67 + i * 0.56
        add_text(slide, label, 0.85, y, 1.7, 0.28, 11, MID)
        rect(slide, 2.62, y + 0.03, 5.35, 0.2, RGBColor(226, 232, 241), radius=True)
        rect(slide, 2.62, y + 0.03, 5.35 * count / max_count, 0.2, color, radius=True)
        add_text(slide, f"{count:,}", 8.08, y - 0.03, 0.75, 0.3, 11, NAVY, True, PP_ALIGN.RIGHT)
    rect(slide, 9.2, 3.32, 3.15, 2.43, PALE_GREEN)
    add_text(slide, "Leakage-safe split", 9.45, 3.55, 2.65, 0.35, 14, GREEN, True)
    add_text(slide, "Train  Sessions 1–3  •  3,259\nValidation  Session 4  •  1,031\nTest  Session 5  •  1,241", 9.45, 4.02, 2.55, 1.02, 12, NAVY)
    add_text(slide, "✓ No speaker overlap", 9.45, 5.17, 2.45, 0.3, 12, GREEN, True)

    # Slide 4
    slide = base_slide(prs, 4, "What has been completed?", "The pipeline is coded and validated; computational experiments have not started")
    completed = [
        "Cleaned 12 GB release to 1.44 GB required audio + labels",
        "Generated metadata and fixed/5-fold session splits",
        "Implemented audio loader, resampling, padding and class mapping",
        "Built CNN and Wav2Vec2/HuBERT attentive-pooling models",
        "Added augmentation, checkpoints, metrics and inference code",
    ]
    rect(slide, 0.75, 1.42, 7.55, 4.95, PALE_GREEN)
    add_text(slide, "COMPLETED", 1.05, 1.68, 2.0, 0.35, 13, GREEN, True)
    for i, item in enumerate(completed):
        y = 2.25 + i * 0.73
        circle(slide, 1.05, y, 0.32, GREEN)
        add_text(slide, "✓", 1.05, y, 0.32, 0.32, 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, item, 1.55, y - 0.04, 6.25, 0.42, 13, NAVY)
    rect(slide, 8.65, 1.42, 3.9, 2.15, PALE_AMBER)
    add_text(slide, "CURRENTLY IN PROGRESS", 8.95, 1.7, 3.25, 0.35, 12, AMBER, True)
    add_text(slide, "Colab environment validation\nand preparation for the first\nCNN baseline run", 8.95, 2.19, 3.1, 0.92, 15, NAVY, True)
    rect(slide, 8.65, 3.92, 3.9, 2.45, LIGHT)
    add_text(slide, "NOT STARTED YET", 8.95, 4.2, 2.8, 0.35, 12, MID, True)
    add_text(slide, "• Model training\n• Hyperparameter tuning\n• Test-set evaluation\n• Result comparison", 8.95, 4.72, 2.85, 1.22, 13, NAVY)

    # Slide 5
    slide = base_slide(prs, 5, "Next steps and expected outputs", "Move from a validated implementation to controlled experiments and final BTP results")
    next_steps = [
        ("1", "Runtime check", "Install dependencies and run one batch on Colab", BLUE),
        ("2", "CNN baseline", "Train, validate and record macro-F1", CYAN),
        ("3", "Main model", "Fine-tune Wav2Vec2 with attentive pooling", AMBER),
        ("4", "Augmentation", "Compare noise, pitch, speed, SpecAugment and Mix-up", GREEN),
        ("5", "Final evaluation", "Five-fold results, confusion matrices and report", RGBColor(130, 108, 190)),
    ]
    for i, (number, title, detail, color) in enumerate(next_steps):
        y = 1.42 + i * 0.91
        circle(slide, 0.85, y, 0.48, color)
        add_text(slide, number, 0.85, y, 0.48, 0.48, 14, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, 1.55, y - 0.02, 2.05, 0.3, 14, NAVY, True)
        add_text(slide, detail, 3.62, y - 0.02, 5.15, 0.35, 12, MID)
        if i < 4:
            rect(slide, 1.075, y + 0.49, 0.03, 0.4, RGBColor(210, 219, 232), radius=False)
    rect(slide, 9.18, 1.48, 3.25, 4.4, NAVY)
    add_text(slide, "FINAL OUTPUT", 9.48, 1.82, 2.65, 0.35, 12, CYAN, True)
    add_text(slide, "A reproducible,\nspeaker-independent\nSER benchmark", 9.48, 2.38, 2.62, 1.35, 23, WHITE, True)
    add_text(slide, "Primary metric\nMACRO-F1", 9.48, 4.2, 2.55, 0.72, 13, RGBColor(197, 210, 229), True)
    add_text(slide, "Target: evidence-backed comparison, not just a single accuracy number", 9.48, 5.13, 2.55, 0.5, 11, CYAN)

    prs.save(OUT)
    print(f"Created {OUT.resolve()} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
