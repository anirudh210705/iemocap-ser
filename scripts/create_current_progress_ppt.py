from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("BTP_Current_Progress_Architectures.pptx")
W, H = 13.333, 7.5

INK = RGBColor(21, 31, 50)
MUTED = RGBColor(91, 105, 125)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(247, 249, 252)
BLUE = RGBColor(42, 104, 214)
CYAN = RGBColor(31, 176, 190)
GREEN = RGBColor(42, 157, 103)
ORANGE = RGBColor(235, 145, 45)
PURPLE = RGBColor(125, 91, 190)
RED = RGBColor(211, 76, 76)
PALE_BLUE = RGBColor(231, 239, 253)
PALE_GREEN = RGBColor(230, 246, 238)
PALE_ORANGE = RGBColor(253, 241, 224)
PALE_PURPLE = RGBColor(240, 234, 250)
GRAY = RGBColor(226, 231, 239)


def textbox(slide, text, x, y, w, h, size=16, color=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    return s


def line(slide, x1, y1, x2, y2, color=MUTED, width=2, arrow=True):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color
    s.line.width = Pt(width)
    if arrow:
        s.line.end_arrowhead = True
    return s


def label_box(slide, title, detail, x, y, w, h, fill, accent, title_size=14, detail_size=10):
    box(slide, x, y, w, h, fill, line=accent)
    box(slide, x, y, 0.08, h, accent, radius=False)
    textbox(slide, title, x + 0.2, y + 0.10, w - 0.3, 0.32, title_size, INK, True)
    textbox(slide, detail, x + 0.2, y + 0.45, w - 0.3, h - 0.52, detail_size, MUTED)


def base(prs, number, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, W, 0.11, BLUE, radius=False)
    textbox(slide, f"{number:02d}", 0.55, 0.34, 0.45, 0.3, 10, BLUE, True)
    textbox(slide, title, 1.05, 0.24, 11.4, 0.55, 26, INK, True)
    if subtitle:
        textbox(slide, subtitle, 1.05, 0.77, 11.3, 0.38, 11, MUTED)
    textbox(slide, "BTP | Speech Emotion Recognition | IEMOCAP", 0.58, 7.13, 5.0, 0.2, 8, MUTED)
    textbox(slide, str(number), 12.35, 7.1, 0.35, 0.2, 8, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def flow(slide, items, y, x0=0.62, total_w=12.05, gap=0.28):
    n = len(items)
    bw = (total_w - gap * (n - 1)) / n
    for i, (title, detail, fill, accent) in enumerate(items):
        x = x0 + i * (bw + gap)
        label_box(slide, title, detail, x, y, bw, 1.12, fill, accent, 13, 9)
        if i < n - 1:
            line(slide, x + bw + 0.03, y + 0.56, x + bw + gap - 0.03, y + 0.56, BLUE, 1.5)


def add_bar_chart(slide, categories, cnn, wav, x, y, w, h):
    data = ChartData()
    data.categories = categories
    data.add_series("CNN", cnn)
    data.add_series("Wav2Vec2", wav)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 80
    chart.value_axis.major_unit = 20
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = CYAN
    chart.series[1].format.fill.solid(); chart.series[1].format.fill.fore_color.rgb = BLUE
    return chart


def build():
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)

    # 1 — title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    box(s, 0, 0, 0.15, H, CYAN, radius=False)
    textbox(s, "BTP CURRENT PROGRESS", 0.85, 0.68, 4.8, 0.35, 13, CYAN, True)
    textbox(s, "Speech Emotion Recognition\n— inside the models", 0.85, 1.27, 7.8, 1.55, 34, INK, True)
    textbox(s, "Two completed 5-fold experiments, their internal architectures,\nand the planned frozen-LLM model.", 0.9, 3.13, 7.2, 0.9, 17, MUTED)
    for i, (n, t, c) in enumerate((("01", "CNN", CYAN), ("02", "Wav2Vec2", BLUE), ("03", "Frozen LLM", PURPLE))):
        x = 0.9 + i * 2.28
        box(s, x, 5.15, 2.0, 0.85, BG, line=c)
        textbox(s, n, x + 0.12, 5.29, 0.4, 0.25, 10, c, True)
        textbox(s, t, x + 0.55, 5.22, 1.25, 0.38, 14, INK, True)
    textbox(s, "14 September 2026", 0.9, 6.62, 2.3, 0.25, 10, MUTED)

    # 2 — metric guide
    s = base(prs, 2, "How to read the evaluation metrics", "Each metric answers a different question; macro-F1 is the main project measure")
    cards = [
        ("Accuracy", "How many predictions were correct overall?\nCorrect ÷ all samples", PALE_BLUE, BLUE),
        ("Precision", "When the model says ‘angry’, how often is it right?\nTP ÷ (TP + FP)", PALE_GREEN, GREEN),
        ("Recall", "Of all truly angry samples, how many were found?\nTP ÷ (TP + FN)", PALE_ORANGE, ORANGE),
        ("F1-score", "Balance of precision and recall. High only when both are high.\n2PR ÷ (P + R)", PALE_PURPLE, PURPLE),
    ]
    for i, item in enumerate(cards):
        x = 0.72 + (i % 2) * 6.12; y = 1.35 + (i // 2) * 1.55
        label_box(s, item[0], item[1], x, y, 5.75, 1.22, item[2], item[3], 16, 12)
    label_box(s, "Macro-F1", "Average the F1 of angry, happy, neutral and sad equally. Best headline metric for imbalanced classes.", 0.72, 4.58, 3.65, 1.15, PALE_PURPLE, PURPLE, 15, 11)
    label_box(s, "Weighted-F1", "Average class F1 after giving larger classes more influence. Reflects the observed class distribution.", 4.61, 4.58, 3.65, 1.15, PALE_BLUE, BLUE, 15, 11)
    label_box(s, "Mean ± SD", "Mean = average of five folds. SD = how much results change across held-out sessions; lower is more stable.", 8.50, 4.58, 3.65, 1.15, PALE_GREEN, GREEN, 15, 11)
    textbox(s, "Interpret together: accuracy shows overall correctness; precision/recall reveal error type; F1 shows their balance.", 0.85, 6.10, 11.55, 0.4, 13, INK, True, PP_ALIGN.CENTER)

    # 3 — CNN internals
    s = base(prs, 3, "Model 1 — CNN baseline: internal architecture", "A compact network learns local time–frequency patterns from log-Mel spectrograms")
    flow(s, [
        ("1. Hear the audio", "Raw speech waveform\n16,000 samples/second", PALE_BLUE, BLUE),
        ("2. Draw a sound map", "64-row log-Mel image\ntime left→right; pitch low→high", PALE_GREEN, GREEN),
        ("3. Find patterns", "3 CNN blocks search for\nlocal voice/emotion clues", PALE_ORANGE, ORANGE),
        ("4. Decide emotion", "compress to 128 numbers\nthen output 4 class scores", PALE_PURPLE, PURPLE),
    ], 1.35, x0=0.55, total_w=12.2, gap=0.24)
    textbox(s, "Inside the pattern finder", 0.72, 2.92, 3.5, 0.35, 16, INK, True)
    flow(s, [
        ("Block 1", "32 filters\nConv + normalize + pool", PALE_ORANGE, ORANGE),
        ("Block 2", "64 filters\nConv + normalize + pool", PALE_ORANGE, ORANGE),
        ("Block 3", "128 filters\nConv + normalize", PALE_ORANGE, ORANGE),
        ("Output", "global average + dropout\nlinear layer → 4 scores", PALE_PURPLE, PURPLE),
    ], 3.40, x0=0.72, total_w=11.85, gap=0.28)
    label_box(s, "Simple interpretation", "The CNN reads the sound map like an image. Early filters detect small changes; deeper filters combine them into emotion-related patterns.", 0.72, 5.08, 7.2, 1.0, WHITE, BLUE, 14, 11)
    label_box(s, "Size", "93,636 trainable parameters", 8.25, 5.08, 4.32, 1.0, PALE_BLUE, CYAN, 14, 16)

    # 4 — CNN results
    s = base(prs, 4, "Model 1 — completed run", "Five-fold test performance; macro-F1 weights all four emotions equally")
    add_bar_chart(s, ["F1", "F2", "F3", "F4", "F5"], [49.59, 42.71, 44.22, 46.30, 43.26], [0, 0, 0, 0, 0], 0.65, 1.35, 7.1, 4.55)
    # hide dummy series by white fill
    chart = s.shapes[-1].chart
    chart.series[1].format.fill.fore_color.rgb = BG
    chart.series[1].format.line.color.rgb = BG
    chart.has_legend = False
    label_box(s, "Accuracy", "45.64%\n± 3.59 pp", 8.10, 1.52, 2.0, 1.35, PALE_BLUE, BLUE, 14, 22)
    label_box(s, "Macro-F1", "45.21%\n± 2.80 pp", 10.35, 1.52, 2.0, 1.35, PALE_GREEN, GREEN, 14, 22)
    label_box(s, "Precision (Fold 1)", "50.27% macro", 8.10, 3.24, 2.0, 1.05, PALE_PURPLE, PURPLE, 12, 16)
    label_box(s, "Recall (Fold 1)", "53.60% macro", 10.35, 3.24, 2.0, 1.05, PALE_ORANGE, ORANGE, 12, 16)
    textbox(s, "Weighted-F1 (5-fold mean): 43.62%\nFold 1 macro-F1: 49.59%", 8.10, 4.57, 4.25, 0.62, 12, MUTED)
    textbox(s, "Interpretation: valid baseline, but substantial emotion confusion remains.", 8.10, 5.38, 4.25, 0.55, 12, INK, True)

    # 5 — Wav2Vec internals
    s = base(prs, 5, "Model 2 — Wav2Vec2: internal architecture", "Pretrained raw-audio representations are adapted while the lower network remains frozen")
    flow(s, [
        ("1. Hear the audio", "Raw speech waveform", PALE_BLUE, BLUE),
        ("2. Make speech features", "7 small CNN layers turn audio\ninto a sequence of feature frames", PALE_GREEN, GREEN),
        ("3. Understand context", "12 Transformer blocks connect\neach moment with the full sentence", PALE_ORANGE, ORANGE),
        ("4. Focus + classify", "attention keeps important moments\nthen outputs 4 emotion scores", PALE_PURPLE, PURPLE),
    ], 1.33, x0=0.55, total_w=12.2, gap=0.24)
    box(s, 0.72, 2.92, 11.9, 1.17, WHITE, line=GRAY)
    textbox(s, "Kept fixed", 0.96, 3.08, 1.0, 0.28, 12, MUTED, True)
    box(s, 2.02, 3.08, 4.1, 0.36, GRAY)
    textbox(s, "feature CNN + Transformer blocks 1–8", 2.10, 3.08, 3.94, 0.36, 11, INK, True, PP_ALIGN.CENTER)
    textbox(s, "Learns", 6.50, 3.08, 0.75, 0.28, 12, GREEN, True)
    box(s, 7.26, 3.08, 4.85, 0.36, PALE_GREEN, line=GREEN)
    textbox(s, "blocks 9–12 + attention + final classifier", 7.38, 3.08, 4.60, 0.36, 11, INK, True, PP_ALIGN.CENTER)
    label_box(s, "What attention does", "It gives every short moment a score, keeps more of the emotionally useful moments, and combines the sequence into one 768-number summary.", 0.72, 4.48, 7.18, 1.25, PALE_PURPLE, PURPLE, 14, 11)
    label_box(s, "Why pretraining helps", "Wav2Vec2 already learned general speech structure from large audio collections; this project only adapts the upper part to emotion.", 8.20, 4.48, 4.42, 1.25, PALE_BLUE, BLUE, 14, 11)

    # 6 — W2V results
    s = base(prs, 6, "Model 2 — completed run", "Wav2Vec2 improves every held-out session, not only the overall average")
    add_bar_chart(s, ["F1", "F2", "F3", "F4", "F5"], [49.59, 42.71, 44.22, 46.30, 43.26], [65.94, 72.51, 64.30, 65.86, 66.56], 0.62, 1.30, 7.5, 4.7)
    label_box(s, "Mean accuracy", "66.67%\n± 2.79 pp", 8.48, 1.48, 3.7, 1.22, PALE_BLUE, BLUE, 14, 22)
    label_box(s, "Mean macro-F1", "67.03%\n± 3.18 pp", 8.48, 2.87, 3.7, 1.12, PALE_GREEN, GREEN, 13, 20)
    label_box(s, "Precision (Fold 1)", "67.14% macro", 8.48, 4.22, 1.72, 1.02, PALE_PURPLE, PURPLE, 11, 14)
    label_box(s, "Recall (Fold 1)", "67.11% macro", 10.46, 4.22, 1.72, 1.02, PALE_ORANGE, ORANGE, 11, 14)
    textbox(s, "Weighted-F1: 66.26%\n+21.82 pp mean macro-F1 vs CNN", 8.48, 5.39, 3.70, 0.58, 12, GREEN, True, PP_ALIGN.CENTER)
    textbox(s, "Best fold: F2 = 72.51%  |  Lowest: F3 = 64.30%", 0.78, 6.18, 7.15, 0.32, 12, MUTED)

    # 7 — progress snapshot
    s = base(prs, 7, "Current progress snapshot", "The acoustic baselines are complete; the multimodal/frozen-LLM experiment is next")
    stages = [
        ("Data pipeline", "Complete", GREEN, "metadata · labels · 5 LOSO splits"),
        ("Model 1 — CNN", "Complete", GREEN, "5/5 folds · report generated"),
        ("Model 2 — Wav2Vec2", "Complete", GREEN, "5/5 folds · report generated"),
        ("Model 3 — frozen LLM", "Coded", ORANGE, "smoke-test/config ready; full run pending"),
        ("Final analysis", "Pending", MUTED, "Model 3 LOSO + comparison + write-up"),
    ]
    for i, (name, status, color, detail) in enumerate(stages):
        y = 1.33 + i * 0.95
        box(s, 0.76, y, 11.75, 0.7, WHITE, line=GRAY)
        box(s, 0.76, y, 0.08, 0.7, color, radius=False)
        textbox(s, name, 1.06, y + 0.11, 3.15, 0.3, 14, INK, True)
        box(s, 4.35, y + 0.15, 1.45, 0.38, color)
        textbox(s, status.upper(), 4.35, y + 0.15, 1.45, 0.38, 10, WHITE, True, PP_ALIGN.CENTER)
        textbox(s, detail, 6.12, y + 0.10, 5.85, 0.34, 12, MUTED)
    box(s, 0.76, 6.17, 11.75, 0.43, PALE_BLUE, line=BLUE)
    textbox(s, "Completed experimental evidence: 10 fold runs = 5 CNN + 5 Wav2Vec2", 0.95, 6.20, 11.35, 0.28, 12, BLUE, True, PP_ALIGN.CENTER)

    # 8 — third model architecture, final slide
    s = base(prs, 8, "Model 3 — planned frozen-LLM architecture", "Implemented in code; full five-fold training has not yet been run")
    flow(s, [
        ("1. Understand speech", "Wav2Vec2 converts audio\nto an emotion-aware summary", PALE_BLUE, BLUE),
        ("2. Translate for LLM", "a small projector changes\n768 values into LLM format", PALE_GREEN, GREEN),
        ("3. Add to prompt", "the audio summary becomes\none special token before the text", PALE_ORANGE, ORANGE),
        ("4. Frozen BLOOMZ", "reads audio token + instruction\nits 1.7B weights do not change", PALE_PURPLE, PURPLE),
        ("5. Choose emotion", "compare only four word scores:\nangry, happy, neutral, sad", PALE_GREEN, GREEN),
    ], 1.35, x0=0.48, total_w=12.38, gap=0.18)
    label_box(s, "What learns during training?", "Upper Wav2Vec2 blocks, attention pooling and the projector. The LLM stays fixed, but the error signal passes through it to teach the audio side.", 0.62, 3.16, 6.12, 1.25, PALE_GREEN, GREEN, 14, 11)
    label_box(s, "What the LLM receives", "[audio token] + “Classify the emotion… Answer with one emotion:”", 6.98, 3.16, 5.73, 1.25, PALE_BLUE, BLUE, 14, 12)
    box(s, 0.62, 4.76, 12.09, 1.0, PALE_ORANGE, line=ORANGE)
    textbox(s, "Main idea", 0.92, 4.94, 1.10, 0.28, 13, ORANGE, True)
    textbox(s, "The projector acts as a bridge: it turns a speech summary into something the language model can process like a token.", 2.16, 4.87, 10.05, 0.45, 13, INK)
    textbox(s, "Next: run the same five LOSO folds and compare with CNN 45.21% and Wav2Vec2 67.03% macro-F1.", 0.82, 6.14, 11.65, 0.38, 14, BLUE, True, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(f"Created {OUT.resolve()} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
